from io import BytesIO
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from zipfile import ZipFile, ZIP_DEFLATED

from django.test import SimpleTestCase, TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from rest_framework.exceptions import ValidationError
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.businesses.models import Business
from apps.clients.models import Client
from apps.core.import_export import read_tabular_file
from apps.core.models import ImportJob
from apps.core.xlsx_reader import read_xlsx_rows
from apps.integrations.models import BusinessEvent


def workbook_bytes(rows=None, transform=None, extra=None):
    workbook = Workbook()
    for row in rows or [["full_name"], ["Synthetic"]]:
        workbook.active.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    if transform is None and extra is None:
        return stream.getvalue()
    output = BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if transform and entry.filename == "xl/worksheets/sheet1.xml":
                content = transform(content.decode()).encode()
            target.writestr(entry.filename, content)
        for name, content in (extra or {}).items():
            target.writestr(name, content)
    return output.getvalue()


class XlsxReaderTests(SimpleTestCase):
    def parse(self, content):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "import.xlsx"
            path.write_bytes(content)
            with patch("apps.core.antivirus.scan_stream", return_value=SimpleNamespace(clean=True)):
                return read_tabular_file(path)

    def test_valid_values_and_underreported_dimensions_preserve_all_rows(self):
        data = workbook_bytes([["full_name", "notes"], ["First", 42], ["Second", " text "]],
                              transform=lambda xml: xml.replace('ref="A1:B3"', 'ref="A1:A1"'),
                              extra={"xl/media/opaque.bin": b"\x89PNG\r\n\x1a\n"})
        self.assertEqual(self.parse(data), [
            {"full_name": "First", "notes": "42"}, {"full_name": "Second", "notes": "text"},
        ])

    @override_settings(IMPORT_MAX_ROWS=2, IMPORT_MAX_COLUMNS=2)
    def test_exact_row_and_column_limits_are_accepted(self):
        self.assertEqual(len(self.parse(workbook_bytes([["full_name", "notes"], ["First", "a"], ["Second", "b"]]))), 2)

    @override_settings(IMPORT_MAX_ROWS=2, IMPORT_MAX_COLUMNS=2)
    def test_dimensions_actual_coordinates_and_counts_are_checked_before_loading(self):
        fixtures = [
            workbook_bytes(transform=lambda xml: xml.replace('ref="A1:A2"', 'ref="A1:XFD1048576"')),
            workbook_bytes([["full_name"], ["First"], ["Second"], ["Third"]],
                           transform=lambda xml: xml.replace('ref="A1:A4"', 'ref="A1:A1"')),
            workbook_bytes([["full_name", "notes", "extra"], ["First", "a", "b"]],
                           transform=lambda xml: xml.replace('ref="A1:C2"', 'ref="A1:A1"')),
            workbook_bytes(transform=lambda xml: xml.replace('r="A2"', 'r="C2"')),
            workbook_bytes(transform=lambda xml: xml.replace('r="2"', 'r="1000000"')),
            workbook_bytes(transform=lambda xml: xml.replace('</sheetData>', '<row/><row/><row/></sheetData>')),
            workbook_bytes(transform=lambda xml: xml.replace('r="A2"', 'r="B2"').replace(
                '</sheetData>', '<row r="3"><c r="B3"><v>1</v></c><c><v>2</v></c></row></sheetData>')),
            workbook_bytes(transform=lambda xml: xml.replace('r="2"', 'r="3"').replace(
                '</sheetData>', '<row/></sheetData>')),
            workbook_bytes(transform=lambda xml: xml.replace('</row>', '<unexpected/></row>')),
        ]
        with patch("openpyxl.load_workbook") as load:
            for data in fixtures:
                with self.subTest(size=len(data)), self.assertRaises(ValidationError):
                    self.parse(data)
            load.assert_not_called()

    def test_archive_limits_and_xml_entities_rejected_before_loading(self):
        fixtures = [
            ({"IMPORT_XLSX_MAX_ENTRIES": 2}, workbook_bytes()),
            ({"IMPORT_XLSX_MAX_EXPANDED_MB": 1}, workbook_bytes(extra={"padding": b"x" * (1024 * 1024)})),
            ({}, workbook_bytes(extra={"padding": b"x" * 100000})),
            ({}, workbook_bytes(transform=lambda xml: '<!DOCTYPE worksheet [<!ENTITY value "expanded">]>' + xml)),
        ]
        with patch("openpyxl.load_workbook") as load:
            for settings, data in fixtures:
                with self.subTest(settings=settings), override_settings(**settings), self.assertRaises(ValidationError):
                    self.parse(data)
            load.assert_not_called()

    def test_empty_malformed_and_invalid_zip_are_validation_errors(self):
        for content in [b"PK\x03\x04bad", workbook_bytes(transform=lambda xml: xml[:-15]),
                        workbook_bytes([["full_name"]]), workbook_bytes([[None]])]:
            with self.subTest(size=len(content)), self.assertRaises(ValidationError):
                self.parse(content)

    def test_relationship_target_with_non_xml_suffix_cannot_bypass_bounds(self):
        data = workbook_bytes(transform=lambda xml: xml.replace('ref="A1:A2"', 'ref="A1:XFD1048576"'))
        output = BytesIO()
        with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename.endswith((".rels", ".xml")):
                    content = content.replace(b"sheet1.xml", b"sheet1.bin")
                target.writestr(entry.filename.replace("sheet1.xml", "sheet1.bin"), content)
        with patch("openpyxl.load_workbook") as load, self.assertRaises(ValidationError):
            self.parse(output.getvalue())
        load.assert_not_called()

    def test_workbook_and_iterator_close_after_conversion_failure_or_early_stop(self):
        for fail in (False, True):
            closed = []
            def rows():
                try:
                    yield ("full_name",)
                    if fail:
                        raise ValueError("invalid workbook value")
                    yield ("Synthetic",)
                    raise AssertionError("Reader consumed beyond caller's requested row")
                finally:
                    closed.append(True)
            book = MagicMock()
            book.active.iter_rows.return_value = rows()
            with self.subTest(fail=fail), patch("openpyxl.load_workbook", return_value=book):
                iterator = read_xlsx_rows(workbook_bytes())
                if fail:
                    with self.assertRaises(ValidationError):
                        next(iterator)
                else:
                    self.assertEqual(next(iterator), {"full_name": "Synthetic"})
                iterator.close()
                book.close.assert_called_once()
                self.assertEqual(closed, [True])


class XlsxImportApiTests(TestCase):
    def setUp(self):
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        setting = override_settings(MEDIA_ROOT=media.name)
        setting.enable()
        self.addCleanup(setting.disable)
        scanner = patch("apps.core.antivirus.scan_stream", return_value=SimpleNamespace(clean=True))
        scanner.start()
        self.addCleanup(scanner.stop)
        self.user = User.objects.create_user(username="xlsx-owner", role=User.Roles.BUSINESS_OWNER)
        self.business = Business.objects.create(owner=self.user, name="XLSX import", slug="xlsx-import")
        self.api = APIClient()
        self.api.force_authenticate(self.user)

    def upload(self, content):
        return self.api.post("/api/import-jobs/", {
            "business": self.business.pk, "entity_type": "clients",
            "source_file": SimpleUploadedFile("clients.xlsx", content,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        }, format="multipart")

    def test_valid_upload_preview_confirm_and_replay(self):
        response = self.upload(workbook_bytes())
        self.assertEqual(response.status_code, 201, response.data)
        url = f'/api/import-jobs/{response.data["id"]}/'
        self.assertEqual(self.api.post(url + "preview/", {}, format="json").status_code, 200)
        self.assertEqual(self.api.post(url + "confirm/").status_code, 200)
        self.assertEqual(self.api.post(url + "confirm/").status_code, 400)
        self.assertEqual(Client.objects.filter(business=self.business).count(), 1)
        self.assertEqual(BusinessEvent.objects.filter(business=self.business, event_type="client.imported").count(), 1)

    def test_bounds_rechecked_at_confirm_without_partial_entities(self):
        response = self.upload(workbook_bytes([["full_name"], ["First"], ["Second"]]))
        self.assertEqual(response.status_code, 201, response.data)
        job = ImportJob.objects.get(pk=response.data["id"])
        with override_settings(IMPORT_MAX_ROWS=1):
            response = self.api.post(f"/api/import-jobs/{job.pk}/confirm/")
        self.assertEqual(response.status_code, 400, response.data)
        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.Statuses.PREVIEWED)
        self.assertFalse(Client.objects.filter(business=self.business).exists())
        self.assertFalse(BusinessEvent.objects.filter(business=self.business).exists())
        self.assertEqual(self.api.post(f"/api/import-jobs/{job.pk}/confirm/").status_code, 200)

    def test_oversized_dimensions_upload_rejected_without_import(self):
        response = self.upload(workbook_bytes(transform=lambda xml: xml.replace('ref="A1:A2"', 'ref="A1:XFD1048576"')))
        self.assertEqual(response.status_code, 400, response.data)
        self.assertFalse(Client.objects.filter(business=self.business).exists())
        self.assertFalse(BusinessEvent.objects.filter(business=self.business).exists())
