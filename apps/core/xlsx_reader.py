"""Bound XLSX expansion and worksheet coordinates before openpyxl allocates rows."""

from io import BytesIO
import re
from xml.parsers import expat
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from django.conf import settings
from rest_framework.exceptions import ValidationError


def _coordinate(value):
    match = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]{0,8})", value)
    if not match:
        raise ValidationError("XLSX contains an invalid cell range.")
    column = 0
    for letter in match[1].upper():
        column = column * 26 + ord(letter) - ord("A") + 1
    if column > settings.IMPORT_MAX_COLUMNS:
        raise ValidationError(f"Import file has too many columns. Maximum is {settings.IMPORT_MAX_COLUMNS}.")
    if int(match[2]) > settings.IMPORT_MAX_ROWS + 1:
        raise ValidationError(f"Import file has too many rows. Maximum is {settings.IMPORT_MAX_ROWS} data rows.")
    return int(match[2]), column


def _check_xml(source):
    parser = expat.ParserCreate(namespace_separator="}")
    worksheet = False
    rows = cells = row_number = column_number = 0
    parents = []

    def start(name, attributes):
        nonlocal worksheet, rows, cells, row_number, column_number
        name = name.rsplit("}", 1)[-1]
        parent = parents[-1] if parents else None
        parents.append(name)
        if name == "worksheet":
            worksheet = True
        if not worksheet:
            return
        if (
            parent == "row" and name != "c"
            or name == "row" and parent != "sheetData"
            or name == "c" and parent != "row"
        ):
            raise ValidationError("XLSX contains an invalid worksheet row.")
        if name == "dimension":
            bounds = attributes.get("ref", "").split(":")
            if len(bounds) not in (1, 2):
                raise ValidationError("XLSX contains an invalid cell range.")
            for bound in bounds:
                _coordinate(bound)
        elif name == "row":
            rows += 1
            cells = 0
            column_number = 0
            row_number, _ = _coordinate("A" + attributes.get("r", str(row_number + 1)))
            if rows > settings.IMPORT_MAX_ROWS + 1:
                raise ValidationError("Import file has too many worksheet rows.")
        elif name == "c":
            cells += 1
            if cells > settings.IMPORT_MAX_COLUMNS:
                raise ValidationError("Import file has too many worksheet columns.")
            if "r" in attributes:
                _, column_number = _coordinate(attributes["r"])
            else:
                column_number += 1
                if column_number > settings.IMPORT_MAX_COLUMNS:
                    raise ValidationError("Import file has too many worksheet columns.")

    def reject_doctype(*args):
        raise ValidationError("XLSX XML document types are not supported.")

    parser.StartElementHandler = start
    parser.EndElementHandler = lambda name: parents.pop()
    parser.StartDoctypeDeclHandler = reject_doctype
    while chunk := source.read(64 * 1024):
        parser.Parse(chunk, False)
    parser.Parse(b"", True)


def validate_xlsx_archive(content):
    with ZipFile(BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > settings.IMPORT_XLSX_MAX_ENTRIES:
            raise ValidationError("XLSX contains too many archive entries.")
        if len({entry.filename for entry in entries}) != len(entries):
            raise ValidationError("XLSX contains duplicate archive entries.")
        if sum(entry.file_size for entry in entries) > settings.IMPORT_XLSX_MAX_EXPANDED_MB * 1024 * 1024:
            raise ValidationError("XLSX expanded content exceeds the import size limit.")
        for entry in entries:
            if entry.flag_bits & 1:
                raise ValidationError("Encrypted XLSX files are not supported.")
            if entry.file_size > max(entry.compress_size, 1) * settings.IMPORT_XLSX_MAX_COMPRESSION_RATIO:
                raise ValidationError("XLSX compression ratio exceeds the import limit.")
        # All workbook XML is checked, including eager metadata/shared strings.
        # Expat callbacks retain no document tree; worksheet limits stop parsing early.
        for entry in entries:
            with archive.open(entry) as source:
                try:
                    _check_xml(source)
                except expat.ExpatError:
                    # Binary images are allowed. XML parts can have arbitrary names
                    # through relationships, so valid XML is checked regardless of suffix.
                    if entry.filename.lower().endswith((".xml", ".rels")):
                        raise


def read_xlsx_rows(content):
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        raise ValidationError("XLSX import requires openpyxl. Install requirements.txt and retry.") from None

    try:
        validate_xlsx_archive(content)
        with BytesIO(content) as source:
            workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
            try:
                sheet = workbook.active
                if sheet is None or not hasattr(sheet, "iter_rows"):
                    raise ValidationError("XLSX has no active worksheet.")
                # Dimensions can under-report actual data. Coordinates were bounded above.
                sheet.reset_dimensions()
                rows = sheet.iter_rows(values_only=True)
                try:
                    headers = [str(value or "").strip() for value in next(rows, ())]
                    if not headers:
                        raise ValidationError("Import file is empty.")
                    for index, values in enumerate(rows, start=1):
                        if index > settings.IMPORT_MAX_ROWS:
                            raise ValidationError("Import file has too many rows.")
                        yield dict(zip(headers, values))
                finally:
                    rows.close()
            finally:
                workbook.close()
    except (BadZipFile, expat.ExpatError, ParseError, KeyError, ValueError, TypeError, IndexError, OSError, NotImplementedError) as exc:
        raise ValidationError("XLSX file is invalid or unsupported.") from exc
